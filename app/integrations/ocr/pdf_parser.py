from typing import List, BinaryIO
import PyPDF2
import openpyxl
import xlrd
import pandas as pd
from docx import Document
import pdf2image
import pytesseract
import io

class DocumentParser:
    """Unified document parser for multiple file formats."""

    SUPPORTED_EXTENSIONS = {'.pdf', '.xlsx', '.xls', '.csv', '.docx', '.html'}

    @staticmethod
    def parse(file_content: BinaryIO, filename: str) -> str:
        """
        Parse document and extract text based on file extension.

        Args:
            file_content: Binary file content
            filename: Name of the file including extension

        Returns:
            Extracted text content

        Raises:
            ValueError: If file type is not supported
        """
        extension = filename.lower().split('.')[-1]
        extension_with_dot = f'.{extension}'

        if extension_with_dot not in DocumentParser.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file type: {extension}. Supported types: {DocumentParser.SUPPORTED_EXTENSIONS}")

        print(f"[INFO] Parsing document: {filename} (type: {extension})")

        try:
            if extension == 'pdf':
                return DocumentParser._parse_pdf(file_content)
            elif extension in ['xlsx', 'xls']:
                return DocumentParser._parse_excel(file_content, extension)
            elif extension == 'csv':
                return DocumentParser._parse_csv(file_content)
            elif extension == 'docx':
                return DocumentParser._parse_docx(file_content)
            elif extension == 'html':
                return DocumentParser._parse_html(file_content)
        except Exception as e:
            print(f"[ERROR] Error parsing {filename}: {str(e)}")
            raise

    @staticmethod
    def _parse_pdf(file_content: BinaryIO) -> str:
        """
        Extract text from PDF file using hybrid approach:
        1. Try native extraction (pdfplumber)
        2. If text is insufficient/empty, fall back to OCR (pdf2image + pytesseract)
        """

        start_pos = file_content.tell()
        file_content.seek(0)
        file_bytes = file_content.read()
        file_content.seek(start_pos)

        file_content_for_plumber = io.BytesIO(file_bytes)

        sample = file_bytes[:500].lower()
        if b"<!doctype html>" in sample or b"<html>" in sample:
            print("[INFO] HTML content detected in PDF file. Routing to HTML parser.")
            file_content_for_plumber.seek(0)
            return DocumentParser._parse_html(file_content_for_plumber)

        try:
            import pdfplumber
            text_parts = []

            with pdfplumber.open(file_content_for_plumber) as pdf:
                total_pages = len(pdf.pages)
                for page in pdf.pages:

                    text = page.extract_text(layout=True)
                    if text and text.strip():
                        text_parts.append(text)

            full_text = '\n\n'.join(text_parts)
        except Exception as e:
            print(f"[ERROR] pdfplumber native extraction failed: {str(e)}")
            full_text = ""
            total_pages = 0

        avg_chars_per_page = len(full_text) / total_pages if total_pages > 0 else 0

        if len(full_text.strip()) < 100 or avg_chars_per_page < 50:
            print(f"[INFO] PDF text extraction insufficient ({len(full_text)} chars, avg {avg_chars_per_page:.1f}/page). Falling back to OCR.")
            try:

                images = pdf2image.convert_from_bytes(file_bytes)
                ocr_text_parts = []

                print(f"[INFO] OCR Processing {len(images)} pages...")
                for i, image in enumerate(images):
                    text = pytesseract.image_to_string(image)
                    if text.strip():
                        ocr_text_parts.append(text)
                    if (i + 1) % 5 == 0:
                        print(f"[INFO] OCR processed page {i+1}/{len(images)}")

                ocr_full_text = '\n\n'.join(ocr_text_parts)

                if len(ocr_full_text.strip()) > len(full_text.strip()):
                    print(f"[INFO] OCR extraction successful. Extracted {len(ocr_full_text)} characters.")
                    return ocr_full_text
                else:
                    print("[WARNING] OCR extracted less text than native. Reverting to native.")
                    return full_text

            except Exception as e:
                print(f"[ERROR] OCR failed: {str(e)}. Returning natively extracted text.")
                return full_text

        print(f"[INFO] Extracted {len(full_text)} characters from PDF ({total_pages} pages) using native extraction")
        return full_text

    @staticmethod
    def _parse_excel(file_content: BinaryIO, extension: str) -> str:
        """Extract text from Excel file (.xlsx or .xls)."""
        text_parts = []

        if extension == 'xlsx':
            workbook = openpyxl.load_workbook(file_content, data_only=True)
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                text_parts.append(f"Sheet: {sheet_name}")

                for row in sheet.iter_rows(values_only=True):
                    row_text = '\t'.join([str(cell) if cell is not None else '' for cell in row])
                    if row_text.strip():
                        text_parts.append(row_text)
        else:
            workbook = xlrd.open_workbook(file_contents=file_content.read())
            for sheet in workbook.sheets():
                text_parts.append(f"Sheet: {sheet.name}")

                for row_idx in range(sheet.nrows):
                    row = sheet.row_values(row_idx)
                    row_text = '\t'.join([str(cell) for cell in row])
                    if row_text.strip():
                        text_parts.append(row_text)

        full_text = '\n'.join(text_parts)
        print(f"[INFO] Extracted {len(full_text)} characters from Excel")
        return full_text

    @staticmethod
    def _parse_csv(file_content: BinaryIO) -> str:
        """Extract text from CSV file."""
        df = pd.read_csv(file_content)

        text_parts = [f"Columns: {', '.join(df.columns.tolist())}"]

        for _, row in df.iterrows():
            row_text = '\t'.join([str(val) for val in row.values])
            text_parts.append(row_text)

        full_text = '\n'.join(text_parts)
        print(f"[INFO] Extracted {len(full_text)} characters from CSV ({len(df)} rows)")
        return full_text

    @staticmethod
    def _parse_docx(file_content: BinaryIO) -> str:
        """Extract text from Word document."""
        doc = Document(file_content)
        text_parts = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        for table in doc.tables:
            for row in table.rows:
                row_text = '\t'.join([cell.text for cell in row.cells])
                if row_text.strip():
                    text_parts.append(row_text)

        full_text = '\n\n'.join(text_parts)
        print(f"[INFO] Extracted {len(full_text)} characters from Word document")
        return full_text

    @staticmethod
    def _parse_html(file_content: BinaryIO) -> str:
        """Extract text from HTML file using regex (safe fallback without BS4)."""
        try:
            raw_content = file_content.read().decode('utf-8')
        except UnicodeDecodeError:
            raw_content = file_content.read().decode('latin-1')

        import re

        text = re.sub(r'<(script|style)[^>]*>.*?</\1>', '', raw_content, flags=re.DOTALL | re.IGNORECASE)

        text = re.sub(r'<[^>]+>', ' ', text)

        text = text.replace('&nbsp;', ' ')
        text = text.replace('&amp;', '&')
        text = text.replace('&lt;', '<')
        text = text.replace('&gt;', '>')

        text = re.sub(r'\s+', ' ', text).strip()

        print(f"[INFO] Extracted {len(text)} characters from HTML document")
        return text
